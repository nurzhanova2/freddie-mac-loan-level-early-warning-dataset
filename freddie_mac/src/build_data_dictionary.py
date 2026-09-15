#!/usr/bin/env python3
"""Export the official Freddie Mac Release 47 file layout to a CSV dictionary."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


SOURCE_FILE = "docs/sources/file_layout_july_2026.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path(SOURCE_FILE))
    parser.add_argument(
        "--output", type=Path,
        default=Path("docs/data_dictionary_release_47.csv"),
    )
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "dataset_component", "field_position", "official_attribute_name",
                "technical_type_format", "max_length", "source_release",
                "predictor_eligibility", "availability_status", "notes",
            ],
        )
        writer.writeheader()
        for worksheet in workbook.worksheets:
            component = (
                "origination" if worksheet.title == "Origination Data File"
                else "monthly_performance"
            )
            for position, name, data_type, max_length in worksheet.iter_rows(
                min_row=3, values_only=True
            ):
                if position is None:
                    continue
                if component == "origination":
                    eligibility, availability = "candidate", "at_origination"
                else:
                    eligibility, availability = "review_required", "monthly_reporting_period"
                writer.writerow({
                    "dataset_component": component,
                    "field_position": position,
                    "official_attribute_name": name,
                    "technical_type_format": data_type,
                    "max_length": max_length,
                    "source_release": "Freddie Mac SFLLD Release 47 (July 2026)",
                    "predictor_eligibility": eligibility,
                    "availability_status": availability,
                    "notes": "Eligibility must be finalised by the leakage register.",
                })


if __name__ == "__main__":
    main()
