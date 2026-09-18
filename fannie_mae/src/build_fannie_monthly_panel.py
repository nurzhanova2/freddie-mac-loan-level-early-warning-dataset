#!/usr/bin/env python3
"""Build a typed Fannie Mae loan-month panel from an official glossary."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

EXPECTED_COLUMNS = 113
PANEL_FIELDS = [
    "loan_identifier", "monthly_reporting_period", "channel", "seller_name",
    "servicer_name", "original_interest_rate", "current_interest_rate", "original_upb",
    "current_actual_upb", "original_loan_term", "origination_date", "first_payment_date",
    "loan_age", "remaining_months_to_legal_maturity", "remaining_months_to_maturity",
    "maturity_date", "original_loan_to_value_ratio_ltv",
    "original_combined_loan_to_value_ratio_cltv", "number_of_borrowers",
    "debt_to_income_dti", "borrower_credit_score_at_origination",
    "co_borrower_credit_score_at_origination", "first_time_home_buyer_indicator",
    "loan_purpose", "property_type", "number_of_units", "occupancy_status",
    "property_state", "metropolitan_statistical_area_msa_or_metropolitan_statistical_division_area_msda",
    "zip_code_short", "mortgage_insurance_percentage", "amortization_type",
    "current_loan_delinquency_status", "loan_payment_history", "modification_flag",
]
EVENT_FIELDS = [
    "loan_identifier", "monthly_reporting_period", "zero_balance_code",
    "zero_balance_effective_date",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def snake_case(value: object) -> str:
    text = str(value).strip().lower().replace("®", "").replace("™", "")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def load_glossary(path: Path) -> tuple[list[dict[str, object]], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Combined Glossary" not in workbook.sheetnames:
        raise ValueError("Expected a worksheet named Combined Glossary.")
    rows = workbook["Combined Glossary"].iter_rows(values_only=True)
    header = next(rows)
    columns = {str(value).strip(): index for index, value in enumerate(header)}
    required = {
        "Field Position", "Field Name", "Description", "Enumerations",
        "Single-Family (SF) Loan Performance", "Type", "Max Length",
    }
    missing = required.difference(columns)
    if missing:
        raise ValueError("Glossary headers missing: {}".format(sorted(missing)))

    entries: dict[int, dict[str, object]] = {}
    for row in rows:
        if row[columns["Field Position"]] is None:
            continue
        position = int(row[columns["Field Position"]])
        name = row[columns["Field Name"]]
        if not name:
            raise ValueError("Glossary position {} has no field name.".format(position))
        entries[position] = {
            "field_position": position,
            "field_name": snake_case(name),
            "official_field_name": str(name).strip(),
            "description": row[columns["Description"]],
            "enumerations": row[columns["Enumerations"]],
            "sf_loan_performance_applicability": row[columns["Single-Family (SF) Loan Performance"]],
            "type": str(row[columns["Type"]]).strip(),
            "max_length": row[columns["Max Length"]],
        }
    expected = set(range(1, EXPECTED_COLUMNS + 1))
    missing_positions = expected.difference(entries)
    if missing_positions:
        raise ValueError("Glossary lacks positions: {}".format(sorted(missing_positions)))
    return [entries[position] for position in sorted(expected)], len(entries)


def schema(columns: list[str], dates: set[str], numerics: set[str]) -> pa.Schema:
    return pa.schema([
        pa.field(
            column,
            pa.timestamp("ns") if column in dates else
            pa.float64() if column in numerics else
            pa.string(),
        )
        for column in columns
    ])


def table(frame: pd.DataFrame, output_schema: pa.Schema) -> pa.Table:
    return pa.Table.from_pandas(frame, schema=output_schema, preserve_index=False)


def write_dictionary(path: Path, fields: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    names = [
        "field_position", "field_name", "official_field_name", "description",
        "enumerations", "sf_loan_performance_applicability", "type", "max_length",
    ]
    with path.open("w", newline="", encoding="utf-8") as target:
        writer = csv.DictWriter(target, fieldnames=names)
        writer.writeheader()
        writer.writerows(fields)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--glossary", required=True, type=Path)
    parser.add_argument("--dictionary", required=True, type=Path)
    parser.add_argument("--cleaned", type=Path, help="Optional full 113-field cleaned parquet; omit for compact analytical processing.")
    parser.add_argument("--panel", required=True, type=Path)
    parser.add_argument("--event-metadata", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    parser.add_argument("--chunksize", type=int, default=50_000)
    args = parser.parse_args()

    fields, all_glossary_positions = load_glossary(args.glossary)
    write_dictionary(args.dictionary, fields)
    field_names = [str(field["field_name"]) for field in fields]
    date_fields = {str(field["field_name"]) for field in fields if field["type"] == "DATE"}
    numeric_fields = {
        str(field["field_name"]) for field in fields if str(field["type"]).strip() == "NUMERIC"
    }
    missing_configured = set(PANEL_FIELDS + EVENT_FIELDS).difference(field_names)
    if missing_configured:
        raise ValueError("Configured fields absent from glossary: {}".format(sorted(missing_configured)))

    output_paths = [args.panel, args.event_metadata, args.report]
    if args.cleaned is not None:
        output_paths.append(args.cleaned)
    for path in output_paths:
        path.parent.mkdir(parents=True, exist_ok=True)
    cleaned_schema = schema(field_names, date_fields, numeric_fields)
    panel_schema = schema(PANEL_FIELDS, date_fields, numeric_fields)
    event_schema = schema(EVENT_FIELDS, date_fields, numeric_fields)
    cleaned_writer = (
        pq.ParquetWriter(args.cleaned, cleaned_schema, compression="zstd")
        if args.cleaned is not None else None
    )
    panel_writer = pq.ParquetWriter(args.panel, panel_schema, compression="zstd")
    event_writer = pq.ParquetWriter(args.event_metadata, event_schema, compression="zstd")
    blank_values = row_count = 0
    invalid_dates: Counter[str] = Counter()
    invalid_numerics: Counter[str] = Counter()
    try:
        reader = pd.read_csv(
            args.source, compression="zip", sep="|", header=None, names=field_names,
            dtype="string", keep_default_na=False, chunksize=args.chunksize, on_bad_lines="error",
        )
        for chunk in reader:
            if len(chunk.columns) != EXPECTED_COLUMNS:
                raise ValueError("Unexpected number of columns during parsing.")
            row_count += len(chunk)
            for column in field_names:
                chunk[column] = chunk[column].str.strip()
            blank_values += int(chunk.eq("").sum().sum())
            chunk = chunk.replace("", pd.NA)
            for column in date_fields:
                before = chunk[column].notna()
                converted = pd.to_datetime(chunk[column], format="%m%Y", errors="coerce")
                invalid_dates[column] += int((before & converted.isna()).sum())
                chunk[column] = converted
            for column in numeric_fields:
                before = chunk[column].notna()
                converted = pd.to_numeric(chunk[column], errors="coerce")
                invalid_numerics[column] += int((before & converted.isna()).sum())
                chunk[column] = converted
            if cleaned_writer is not None:
                cleaned_writer.write_table(table(chunk, cleaned_schema))
            panel_writer.write_table(table(chunk[PANEL_FIELDS], panel_schema))
            event_writer.write_table(table(chunk[EVENT_FIELDS], event_schema))
    finally:
        if cleaned_writer is not None:
            cleaned_writer.close()
        panel_writer.close()
        event_writer.close()

    report = {
        "rows_written": row_count,
        "raw_column_count": EXPECTED_COLUMNS,
        "official_glossary_path": args.glossary.as_posix(),
        "official_glossary_sha256": sha256(args.glossary),
        "official_glossary_position_count": all_glossary_positions,
        "mapped_raw_positions": [1, EXPECTED_COLUMNS],
        "glossary_position_not_present_in_raw_archive": [114] if all_glossary_positions > EXPECTED_COLUMNS else [],
        "dictionary_path": args.dictionary.as_posix(),
        "full_cleaned_parquet_written": args.cleaned is not None,
        "blank_values_converted_to_null": blank_values,
        "cleaning_rule": "Whitespace and blank strings converted to null; official field names and types imported from the Fannie Mae glossary; special codes are not recoded.",
        "invalid_dates_converted_to_null": dict(sorted(invalid_dates.items())),
        "invalid_numeric_values_converted_to_null": dict(sorted(invalid_numerics.items())),
        "panel_key": ["loan_identifier", "monthly_reporting_period"],
        "panel_feature_policy": "Candidate predictors and current performance context retained; zero-balance fields are event metadata, not model features.",
        "unmapped_field_positions": [],
        "status": "panel_created_with_official_glossary_mapping_positions_1_to_113",
    }
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
